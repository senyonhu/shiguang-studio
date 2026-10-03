"""Extract source records and provenance; never infer or execute story content."""

from __future__ import annotations

import csv
import hashlib
import io
import math
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Any


SUPPORTED_FORMATS = {".xlsx", ".csv", ".tsv", ".txt", ".md", ".docx"}
MAX_FILE_BYTES = 50 * 1024 * 1024
MAX_EXPANDED_BYTES = 250 * 1024 * 1024
MAX_TABLE_CELLS = 250_000


class SourceError(ValueError):
    def __init__(self, code: str, source_path: str, message: str):
        self.code = code
        self.source_path = source_path
        super().__init__(message)


def _warning(code: str, message: str, **location: Any) -> dict[str, Any]:
    return {"code": code, "message": message, **location}


def _json_value(value: Any) -> str | int | float | bool | None:
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("表格包含非有限数字")
        return value
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    return str(value)


def _check_archive(data: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
            raise ValueError("文档解压大小超过提取上限")


def _decode(data: bytes, warnings: list[dict[str, Any]]) -> tuple[str, str]:
    try:
        return data.decode("utf-8-sig"), "utf-8-sig"
    except UnicodeDecodeError:
        text = data.decode("gb18030")
        warnings.append(_warning("ENCODING_FALLBACK", "源文件采用 gb18030 解码；原文件字节和哈希保持不变", encoding="gb18030"))
        return text, "gb18030"


def _extract_xlsx(data: bytes, warnings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    _check_archive(data)
    formulas = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    cached = load_workbook(io.BytesIO(data), data_only=True, keep_links=False)
    records: list[dict[str, Any]] = []
    try:
        if sum(sheet.max_row * sheet.max_column for sheet in formulas) > MAX_TABLE_CELLS:
            raise ValueError("工作表单元格范围超过提取上限；请拆分资料或清除多余格式区域")
        for sheet in formulas:
            merged_ranges = list(sheet.merged_cells.ranges)
            merged_lookup: dict[tuple[int, int], str] = {}
            if sum((item.max_row - item.min_row + 1) * (item.max_col - item.min_col + 1) for item in merged_ranges) > MAX_TABLE_CELLS:
                raise ValueError("合并单元格覆盖范围超过提取上限")
            for merged in merged_ranges:
                records.append({
                    "kind": "merged_range", "sheet": sheet.title,
                    "range": str(merged), "anchor": sheet.cell(merged.min_row, merged.min_col).coordinate,
                })
                for position in merged.cells:
                    merged_lookup[position] = str(merged)
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    merged = merged_lookup.get((cell.row, cell.column))
                    formula = cell.value if cell.data_type == "f" else None
                    value = cached[sheet.title][cell.coordinate].value if formula is not None else cell.value
                    records.append({
                        "kind": "cell", "sheet": sheet.title, "coordinate": cell.coordinate,
                        "row": cell.row, "column": cell.column, "value": _json_value(value),
                        "formula": formula, "merged_range": merged,
                    })
                    if formula is not None and value is None:
                        warnings.append(_warning("FORMULA_WITHOUT_CACHE", "公式没有可用缓存值；没有执行或猜算公式", sheet=sheet.title, coordinate=cell.coordinate, formula=formula))
            if getattr(sheet, "_images", None):
                warnings.append(_warning("EMBEDDED_IMAGES_NOT_EXTRACTED", "工作表包含嵌入图片；此提取器只提取表格内容，需另行登记图片", sheet=sheet.title))
    finally:
        formulas.close()
        cached.close()
    return records


def _extract_delimited(text: str, extension: str) -> list[dict[str, Any]]:
    delimiter = "\t" if extension == ".tsv" else ","
    if extension == ".csv":
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
        except csv.Error:
            pass
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    records = []
    previous_line = 0
    cells = 0
    for index, row in enumerate(reader, 1):
        cells += len(row)
        if cells > MAX_TABLE_CELLS:
            raise ValueError("表格单元格超过提取上限")
        records.append({
            "kind": "row", "row": index, "line_start": previous_line + 1,
            "line_end": reader.line_num, "values": row, "delimiter": delimiter,
        })
        previous_line = reader.line_num
    return records


def _extract_paragraphs(text: str) -> list[dict[str, Any]]:
    records = []
    paragraph: list[str] = []
    start = 1
    for number, line in enumerate(text.splitlines(), 1):
        if line.strip():
            if not paragraph:
                start = number
            paragraph.append(line)
        elif paragraph:
            records.append({"kind": "paragraph", "paragraph": len(records) + 1, "line_start": start, "line_end": number - 1, "text": "\n".join(paragraph)})
            paragraph = []
    if paragraph:
        records.append({"kind": "paragraph", "paragraph": len(records) + 1, "line_start": start, "line_end": start + len(paragraph) - 1, "text": "\n".join(paragraph)})
    return records


def _extract_docx(data: bytes, warnings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    _check_archive(data)
    document = Document(io.BytesIO(data))
    records = []
    paragraph_index = 0
    table_index = 0
    cell_count = 0
    for block_index, element in enumerate(document.element.body, 1):
        if element.tag == qn("w:p"):
            paragraph_index += 1
            paragraph = Paragraph(element, document._body)
            records.append({
                "kind": "paragraph", "block": block_index, "paragraph": paragraph_index,
                "text": paragraph.text, "style": paragraph.style.name if paragraph.style else None,
            })
        elif element.tag == qn("w:tbl"):
            table_index += 1
            table = Table(element, document._body)
            for row_index, row in enumerate(table.rows, 1):
                cell_count += len(row.cells)
                if cell_count > MAX_TABLE_CELLS:
                    raise ValueError("文档表格单元格超过提取上限")
                records.append({
                    "kind": "table_row", "block": block_index, "table": table_index,
                    "row": row_index, "values": [cell.text for cell in row.cells],
                })
            if element.findall(".//" + qn("w:tbl")):
                warnings.append(_warning("NESTED_TABLE_PARTIAL", "嵌套表格未单独提取；普通表格行仍保留", table=table_index))
    if document.inline_shapes:
        warnings.append(_warning("EMBEDDED_IMAGES_NOT_EXTRACTED", "DOCX 包含嵌入媒体；正文提取不登记这些素材"))
    if document.element.findall(".//" + qn("w:txbxContent")):
        warnings.append(_warning("TEXTBOX_NOT_EXTRACTED", "DOCX 文本框不属于普通正文段落，尚未提取"))
    return records


def extract_file(path: Path) -> dict[str, Any]:
    """Return JSON-friendly source records, hash and explicit format limits.

    Caller authorization/root checks precede this read. A single byte snapshot
    supplies both the hash and parsers so a concurrent file change cannot make
    the reported hash refer to different content than the extracted records.
    """
    path = Path(path)
    source_path = str(path.resolve())
    extension = path.suffix.lower()
    if extension not in SUPPORTED_FORMATS:
        raise SourceError("UNSUPPORTED_FORMAT", source_path, f"暂不支持 {extension or '无扩展名'}；支持 xlsx/csv/tsv/txt/md/docx")
    try:
        with path.open("rb") as stream:
            data = stream.read(MAX_FILE_BYTES + 1)
        if len(data) > MAX_FILE_BYTES:
            raise SourceError("SOURCE_TOO_LARGE", source_path, "源文件超过 50 MiB 提取上限")
    except OSError as error:
        raise SourceError("SOURCE_UNREADABLE", source_path, f"不能读取源文件: {error}") from error
    warnings: list[dict[str, Any]] = []
    result: dict[str, Any] = {
        "format": extension[1:], "source_path": source_path,
        "hash": hashlib.sha256(data).hexdigest(), "size_bytes": len(data),
        "records": [], "warnings": warnings,
    }
    try:
        if extension == ".xlsx":
            result["records"] = _extract_xlsx(data, warnings)
        elif extension == ".docx":
            result["records"] = _extract_docx(data, warnings)
        else:
            text, encoding = _decode(data, warnings)
            if "\x00" in text:
                raise ValueError("文本包含 NUL 字节，可能不是支持的文本资料")
            result["encoding"] = encoding
            result["records"] = _extract_delimited(text, extension) if extension in (".csv", ".tsv") else _extract_paragraphs(text)
    except (ImportError, ModuleNotFoundError) as error:
        raise SourceError("EXTRACTOR_UNAVAILABLE", source_path, f"缺少提取库: {error}") from error
    except Exception as error:
        raise SourceError("SOURCE_INVALID", source_path, f"源文件损坏、格式不符或超过边界: {error}") from error
    return result
