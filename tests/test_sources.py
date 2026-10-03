import hashlib
import io
import zipfile
from datetime import date

import pytest
from docx import Document
from openpyxl import Workbook

from renpy_bridge.sources import SourceError, extract_file


def test_xlsx_coordinates_merged_cells_formulas_dates_and_hash(tmp_path):
    path = tmp_path / "第一幕 资料.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "第一幕"
    sheet["A1"] = "幕"
    sheet["A2"] = "第一幕"
    sheet.merge_cells("A2:A3")
    sheet["B2"] = "保留 [原文] {标签}"
    sheet["C2"] = "=1+1"
    sheet["D2"] = date(2026, 10, 2)
    workbook.create_sheet("第二幕")["A1"] = "第1帧"
    workbook.save(path)
    result = extract_file(path)
    cells = {(record["sheet"], record["coordinate"]): record for record in result["records"] if record["kind"] == "cell"}
    assert result["hash"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert result["format"] == "xlsx" and result["source_path"] == str(path.resolve())
    assert cells[("第一幕", "A2")]["merged_range"] == "A2:A3"
    assert cells[("第一幕", "B2")]["value"] == "保留 [原文] {标签}"
    assert cells[("第一幕", "C2")]["formula"] == "=1+1"
    assert cells[("第一幕", "C2")]["value"] is None
    assert cells[("第一幕", "D2")]["value"].startswith("2026-10-02")
    assert cells[("第二幕", "A1")]["value"] == "第1帧"
    assert any(record["kind"] == "merged_range" and record["anchor"] == "A2" for record in result["records"])
    assert any(warning["code"] == "FORMULA_WITHOUT_CACHE" and warning["coordinate"] == "C2" for warning in result["warnings"])


def test_xlsx_uses_existing_formula_cache_without_executing_formula(tmp_path):
    workbook = Workbook()
    workbook.active["A1"] = "=1+1"
    original = io.BytesIO()
    workbook.save(original)
    modified = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original.getvalue())) as source, zipfile.ZipFile(modified, "w") as output:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                assert b"<f>1+1</f><v></v>" in content
                content = content.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>2</v>")
            output.writestr(item, content)
    path = tmp_path / "cached.xlsx"
    path.write_bytes(modified.getvalue())
    result = extract_file(path)
    cell = next(record for record in result["records"] if record["kind"] == "cell")
    assert cell["value"] == 2 and cell["formula"] == "=1+1"
    assert not result["warnings"]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "gb18030"])
def test_csv_chinese_encodings_multiline_and_bom(tmp_path, encoding):
    path = tmp_path / "对白.csv"
    text = '幕,帧,台词\n第一幕,1,"你好\n世界"\n'
    path.write_bytes(text.encode(encoding))
    result = extract_file(path)
    assert result["records"][0]["values"] == ["幕", "帧", "台词"]
    assert result["records"][1]["values"] == ["第一幕", "1", "你好\n世界"]
    assert result["records"][1]["line_start"] == 2 and result["records"][1]["line_end"] == 3
    assert bool(result["warnings"]) == (encoding == "gb18030")


def test_tsv_keeps_tabs_as_columns(tmp_path):
    path = tmp_path / "归属.tsv"
    path.write_text("幕\t帧\t人物\n一\t1\t安娜\n", encoding="utf-8")
    result = extract_file(path)
    assert result["format"] == "tsv"
    assert result["records"][1]["values"] == ["一", "1", "安娜"]


def test_semicolon_csv_can_be_extracted(tmp_path):
    path = tmp_path / "semicolon.csv"
    path.write_text("幕;帧;人物\n一;1;安娜\n", encoding="utf-8")
    result = extract_file(path)
    assert result["records"][0]["delimiter"] == ";"
    assert result["records"][1]["values"] == ["一", "1", "安娜"]


@pytest.mark.parametrize("extension", ["txt", "md"])
def test_text_paragraphs_preserve_literal_content_and_line_ranges(tmp_path, extension):
    path = tmp_path / f"资料.{extension}"
    path.write_text("# 第一幕\n  第一帧 [原文] {不执行}\n\n第二帧\n继续。\n", encoding="utf-8")
    records = extract_file(path)["records"]
    assert records[0]["text"] == "# 第一幕\n  第一帧 [原文] {不执行}"
    assert (records[0]["line_start"], records[0]["line_end"]) == (1, 2)
    assert records[1]["text"] == "第二帧\n继续。"
    assert (records[1]["line_start"], records[1]["line_end"]) == (4, 5)


def test_docx_body_order_and_basic_tables_are_preserved(tmp_path):
    path = tmp_path / "说明.docx"
    document = Document()
    document.add_paragraph("第一幕说明")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "帧"
    table.cell(0, 1).text = "台词"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "你好 [原文]"
    document.add_paragraph("第二幕说明")
    document.save(path)
    result = extract_file(path)
    assert [record["kind"] for record in result["records"]] == ["paragraph", "table_row", "table_row", "paragraph"]
    assert result["records"][2]["values"] == ["1", "你好 [原文]"]
    assert result["records"][0]["block"] < result["records"][1]["block"] < result["records"][-1]["block"]


@pytest.mark.parametrize("extension", ["xlsx", "docx"])
def test_corrupt_packaged_documents_raise_readable_errors(tmp_path, extension):
    path = tmp_path / f"坏文件.{extension}"
    path.write_bytes(b"not a packaged document")
    with pytest.raises(SourceError) as captured:
        extract_file(path)
    assert captured.value.code == "SOURCE_INVALID"
    assert captured.value.source_path == str(path.resolve())


def test_unreadable_and_unsupported_formats_are_distinct(tmp_path):
    with pytest.raises(SourceError) as missing:
        extract_file(tmp_path / "missing.txt")
    assert missing.value.code == "SOURCE_UNREADABLE"
    for extension in ("xls", "pdf", "exe"):
        with pytest.raises(SourceError) as unsupported:
            extract_file(tmp_path / f"资料.{extension}")
        assert unsupported.value.code == "UNSUPPORTED_FORMAT"


def test_invalid_text_and_unterminated_csv_are_not_guessed(tmp_path):
    text = tmp_path / "binary.txt"
    text.write_bytes(b"hello\x00world")
    with pytest.raises(SourceError, match="NUL"):
        extract_file(text)
    csv_path = tmp_path / "broken.csv"
    csv_path.write_text('a,b\n1,"unterminated', encoding="utf-8")
    with pytest.raises(SourceError):
        extract_file(csv_path)


def test_text_like_commands_are_only_data(tmp_path):
    path = tmp_path / "commands.txt"
    raw = "__import__('os').system('echo unwanted')\n=HYPERLINK(\"https://example.test\")"
    path.write_text(raw, encoding="utf-8")
    assert extract_file(path)["records"][0]["text"] == raw
